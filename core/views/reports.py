from django.shortcuts import render
from django.contrib.auth.decorators import login_required
from core.permissions import staff_required, admin_required
from django.contrib import messages
from django.core.paginator import Paginator
from django.db.models import Sum, Count, Q
from django.http import HttpResponse
from datetime import date, datetime, timedelta
from decimal import Decimal
import calendar

from core.models import (
    Stall, Tenant, RentalContract, Billing, Payment,
    TenantLedger, Notice, AuditLog, MarketSection, UserProfile
)


def get_date_range(request):
    date_from = request.GET.get('date_from')
    date_to = request.GET.get('date_to')
    if not date_from:
        date_from = date.today().replace(day=1).isoformat()
    if not date_to:
        date_to = date.today().isoformat()
    return date_from, date_to


def get_month_year(request):
    month = request.GET.get('month', date.today().month)
    year = request.GET.get('year', date.today().year)
    try:
        month = int(month)
        year = int(year)
    except (ValueError, TypeError):
        month = date.today().month
        year = date.today().year
    return month, year


@staff_required
def report_dashboard(request):
    today = date.today()
    month, year = get_month_year(request)
    _, last_day = calendar.monthrange(year, month)

    total_stalls = Stall.objects.count()
    occupied = Stall.objects.filter(status='Occupied').count()
    vacant = Stall.objects.filter(status='Vacant').count()

    billings = Billing.objects.filter(billing_month=month, billing_year=year)
    total_billed = billings.aggregate(total=Sum('total_due'))['total'] or 0
    total_collected = billings.aggregate(total=Sum('amount_paid'))['total'] or 0
    total_balance = billings.aggregate(total=Sum('balance'))['total'] or 0
    paid_count = billings.filter(status='Paid').count()
    unpaid_count = billings.filter(status='Unpaid').count()
    overdue_count = billings.filter(status='Overdue').count()
    total_count = billings.count()

    collection_rate = round(float(total_collected) / float(total_billed) * 100, 2) if total_billed else 0

    context = {
        'total_stalls': total_stalls,
        'occupied': occupied,
        'vacant': vacant,
        'month': month,
        'year': year,
        'total_billed': total_billed,
        'total_collected': total_collected,
        'total_balance': total_balance,
        'paid_count': paid_count,
        'unpaid_count': unpaid_count,
        'overdue_count': overdue_count,
        'total_count': total_count,
        'collection_rate': collection_rate,
        'month_name': calendar.month_name[month],
    }
    return render(request, 'reports/dashboard.html', context)


@staff_required
def report_daily(request):
    date_from, date_to = get_date_range(request)
    payments = Payment.objects.filter(
        status__in=['Paid', 'Partial'],
        payment_date__gte=date_from,
        payment_date__lte=date_to,
    ).select_related('tenant', 'stall', 'collected_by').order_by('payment_date')

    paginator = Paginator(payments, 50)
    page_obj = paginator.get_page(request.GET.get('page'))

    total_collected = payments.aggregate(total=Sum('amount_paid'))['total'] or 0
    total_count = payments.count()

    context = {
        'page_obj': page_obj,
        'date_from': date_from,
        'date_to': date_to,
        'total_collected': total_collected,
        'total_count': total_count,
        'report_type': 'Daily Collection Report',
    }
    return render(request, 'reports/daily.html', context)


@staff_required
def report_weekly(request):
    today = date.today()
    weekday = today.weekday()
    start_of_week = today - timedelta(days=weekday)
    end_of_week = start_of_week + timedelta(days=6)

    date_from = request.GET.get('date_from', start_of_week.isoformat())
    date_to = request.GET.get('date_to', end_of_week.isoformat())

    payments = Payment.objects.filter(
        status__in=['Paid', 'Partial'],
        payment_date__gte=date_from,
        payment_date__lte=date_to,
    ).select_related('tenant', 'stall', 'collected_by').order_by('payment_date')

    paginator = Paginator(payments, 50)
    page_obj = paginator.get_page(request.GET.get('page'))

    total_collected = payments.aggregate(total=Sum('amount_paid'))['total'] or 0
    total_count = payments.count()

    context = {
        'page_obj': page_obj,
        'date_from': date_from,
        'date_to': date_to,
        'total_collected': total_collected,
        'total_count': total_count,
        'report_type': 'Weekly Collection Report',
    }
    return render(request, 'reports/weekly.html', context)


@staff_required
def report_monthly(request):
    month, year = get_month_year(request)
    _, last_day = calendar.monthrange(year, month)
    start_date = date(year, month, 1)
    end_date = date(year, month, last_day)

    payments = Payment.objects.filter(
        status__in=['Paid', 'Partial'],
        payment_date__gte=start_date,
        payment_date__lte=end_date,
    ).select_related('tenant', 'stall', 'collected_by').order_by('payment_date')

    paginator = Paginator(payments, 50)
    page_obj = paginator.get_page(request.GET.get('page'))

    total_collected = payments.aggregate(total=Sum('amount_paid'))['total'] or 0
    total_count = payments.count()

    billings = Billing.objects.filter(billing_month=month, billing_year=year)
    total_billed = billings.aggregate(total=Sum('total_due'))['total'] or 0

    context = {
        'page_obj': page_obj,
        'month': month,
        'year': year,
        'month_name': calendar.month_name[month],
        'total_collected': total_collected,
        'total_billed': total_billed,
        'total_count': total_count,
        'report_type': 'Monthly Collection Report',
    }
    return render(request, 'reports/monthly.html', context)


@staff_required
def report_annual(request):
    year = request.GET.get('year', date.today().year)
    try:
        year = int(year)
    except (ValueError, TypeError):
        year = date.today().year

    monthly_data = []
    for m in range(1, 13):
        payments = Payment.objects.filter(
            status__in=['Paid', 'Partial'],
            payment_date__year=year,
            payment_date__month=m,
        )
        total = payments.aggregate(total=Sum('amount_paid'))['total'] or 0
        count = payments.count()
        monthly_data.append({
            'month': m,
            'month_name': calendar.month_name[m],
            'total': total,
            'count': count,
        })

    grand_total = sum(d['total'] for d in monthly_data)
    grand_count = sum(d['count'] for d in monthly_data)

    context = {
        'year': year,
        'monthly_data': monthly_data,
        'grand_total': grand_total,
        'grand_count': grand_count,
        'report_type': 'Annual Collection Report',
    }
    return render(request, 'reports/annual.html', context)


@staff_required
def report_section(request):
    month, year = get_month_year(request)

    sections = MarketSection.objects.annotate(
        stall_count=Count('stalls')
    )
    section_data = []
    for section in sections:
        section_stalls = Stall.objects.filter(section=section)
        occupied = section_stalls.filter(status='Occupied').count()
        vacant = section_stalls.filter(status='Vacant').count()

        billings = Billing.objects.filter(
            stall__section=section,
            billing_month=month,
            billing_year=year,
        )
        total_billed = billings.aggregate(total=Sum('total_due'))['total'] or 0
        total_collected = billings.aggregate(total=Sum('amount_paid'))['total'] or 0
        total_balance = billings.aggregate(total=Sum('balance'))['total'] or 0

        section_data.append({
            'section': section,
            'total_stalls': section_stalls.count(),
            'occupied': occupied,
            'vacant': vacant,
            'total_billed': total_billed,
            'total_collected': total_collected,
            'total_balance': total_balance,
        })

    context = {
        'section_data': section_data,
        'month': month,
        'year': year,
        'month_name': calendar.month_name[month],
        'report_type': 'Collection by Section',
    }
    return render(request, 'reports/section.html', context)


@staff_required
def report_collector(request):
    date_from, date_to = get_date_range(request)

    collectors = UserProfile.objects.filter(
        role__in=['collector', 'cashier', 'admin']
    ).select_related('user')
    collector_data = []
    for profile in collectors:
        payments = Payment.objects.filter(
            collected_by=profile.user,
            status__in=['Paid', 'Partial'],
            payment_date__gte=date_from,
            payment_date__lte=date_to,
        )
        total = payments.aggregate(total=Sum('amount_paid'))['total'] or 0
        count = payments.count()
        collector_data.append({
            'collector': profile.user,
            'total': total,
            'count': count,
        })

    grand_total = sum(d['total'] for d in collector_data)
    grand_count = sum(d['count'] for d in collector_data)

    context = {
        'collector_data': collector_data,
        'date_from': date_from,
        'date_to': date_to,
        'grand_total': grand_total,
        'grand_count': grand_count,
        'report_type': 'Collection by Collector',
    }
    return render(request, 'reports/collector.html', context)


@staff_required
def report_paid_tenants(request):
    month, year = get_month_year(request)
    billings = Billing.objects.filter(
        billing_month=month,
        billing_year=year,
        status='Paid',
    ).select_related('tenant', 'stall').order_by('tenant__full_name')

    paginator = Paginator(billings, 50)
    page_obj = paginator.get_page(request.GET.get('page'))

    total_collected = billings.aggregate(total=Sum('amount_paid'))['total'] or 0
    total_count = billings.count()

    context = {
        'page_obj': page_obj,
        'month': month,
        'year': year,
        'month_name': calendar.month_name[month],
        'total_collected': total_collected,
        'total_count': total_count,
        'report_type': 'Paid Tenants Report',
    }
    return render(request, 'reports/paid_tenants.html', context)


@staff_required
def report_unpaid_tenants(request):
    month, year = get_month_year(request)
    billings = Billing.objects.filter(
        billing_month=month,
        billing_year=year,
    ).filter(
        Q(status='Unpaid') | Q(status='Partial') | Q(status='Overdue')
    ).select_related('tenant', 'stall').order_by('-balance')

    paginator = Paginator(billings, 50)
    page_obj = paginator.get_page(request.GET.get('page'))

    total_unpaid = billings.aggregate(total=Sum('balance'))['total'] or 0
    total_count = billings.count()

    context = {
        'page_obj': page_obj,
        'month': month,
        'year': year,
        'month_name': calendar.month_name[month],
        'total_unpaid': total_unpaid,
        'total_count': total_count,
        'report_type': 'Unpaid Tenants Report',
    }
    return render(request, 'reports/unpaid_tenants.html', context)


@staff_required
def report_overdue(request):
    billings = Billing.objects.filter(
        status='Overdue'
    ).select_related('tenant', 'stall').order_by('-balance')

    paginator = Paginator(billings, 50)
    page_obj = paginator.get_page(request.GET.get('page'))

    total_overdue = billings.aggregate(total=Sum('balance'))['total'] or 0
    total_count = billings.count()

    context = {
        'page_obj': page_obj,
        'total_overdue': total_overdue,
        'total_count': total_count,
        'report_type': 'Overdue Accounts Report',
    }
    return render(request, 'reports/overdue.html', context)


@staff_required
def report_ledger(request):
    tenant_id = request.GET.get('tenant')
    date_from, date_to = get_date_range(request)

    entries = TenantLedger.objects.select_related('tenant', 'billing', 'payment').order_by('transaction_date')
    if tenant_id:
        entries = entries.filter(tenant_id=tenant_id)
    if date_from:
        entries = entries.filter(transaction_date__gte=date_from)
    if date_to:
        entries = entries.filter(transaction_date__lte=date_to)

    paginator = Paginator(entries, 50)
    page_obj = paginator.get_page(request.GET.get('page'))

    total_debit = entries.aggregate(total=Sum('debit'))['total'] or 0
    total_credit = entries.aggregate(total=Sum('credit'))['total'] or 0

    context = {
        'page_obj': page_obj,
        'tenant_id': tenant_id,
        'date_from': date_from,
        'date_to': date_to,
        'total_debit': total_debit,
        'total_credit': total_credit,
        'report_type': 'Tenant Ledger Report',
    }
    return render(request, 'reports/ledger.html', context)


@staff_required
def report_occupancy(request):
    stalls = Stall.objects.select_related('section', 'stall_type').all().order_by('section__name', 'stall_number')

    section_filter = request.GET.get('section')
    status_filter = request.GET.get('status')

    if section_filter:
        stalls = stalls.filter(section_id=section_filter)
    if status_filter:
        stalls = stalls.filter(status=status_filter)

    paginator = Paginator(stalls, 50)
    page_obj = paginator.get_page(request.GET.get('page'))

    total = stalls.count()
    occupied = stalls.filter(status='Occupied').count()
    vacant = stalls.filter(status='Vacant').count()

    sections = MarketSection.objects.all().order_by('name')

    context = {
        'page_obj': page_obj,
        'total': total,
        'occupied': occupied,
        'vacant': vacant,
        'sections': sections,
        'section_filter': section_filter,
        'status_filter': status_filter,
        'report_type': 'Stall Occupancy Report',
    }
    return render(request, 'reports/occupancy.html', context)


@staff_required
def report_vacant(request):
    stalls = Stall.objects.filter(status='Vacant').select_related('section', 'stall_type').order_by('stall_number')
    section_filter = request.GET.get('section')
    if section_filter:
        stalls = stalls.filter(section_id=section_filter)

    paginator = Paginator(stalls, 50)
    page_obj = paginator.get_page(request.GET.get('page'))

    context = {
        'page_obj': page_obj,
        'total': stalls.count(),
        'section_filter': section_filter,
        'sections': MarketSection.objects.all().order_by('name'),
        'report_type': 'Vacant Stall Report',
    }
    return render(request, 'reports/vacant.html', context)


@staff_required
def report_export_pdf(request):
    report_type = request.GET.get('report', 'monthly')
    month, year = get_month_year(request)
    date_from, date_to = get_date_range(request)

    try:
        import weasyprint
    except (ImportError, OSError) as e:
        messages.error(request, f'PDF export unavailable on this system ({e}). Install weasyprint system dependencies or use Excel/Print.')
        return render(request, 'reports/export_pdf.html', {'error': True})

    # Route to correct dataset and template
    if report_type == 'monthly':
        billings = Billing.objects.filter(billing_month=month, billing_year=year).select_related('tenant', 'stall')
        total_billed = billings.aggregate(total=Sum('total_due'))['total'] or 0
        total_collected = billings.aggregate(total=Sum('amount_paid'))['total'] or 0
        template = 'reports/pdf/monthly.html'
        context = {'billings': billings, 'month': month, 'year': year, 'month_name': __import__('calendar').month_name[month], 'total_billed': total_billed, 'total_collected': total_collected, 'report_type': 'Monthly Collection Report'}
    elif report_type == 'daily':
        payments = Payment.objects.filter(status__in=['Paid', 'Partial'], payment_date__gte=date_from, payment_date__lte=date_to).select_related('tenant', 'stall', 'collected_by')
        total_collected = payments.aggregate(total=Sum('amount_paid'))['total'] or 0
        template = 'reports/pdf/daily.html'
        context = {'payments': payments, 'date_from': date_from, 'date_to': date_to, 'total_collected': total_collected, 'report_type': 'Daily Collection Report'}
    elif report_type == 'weekly':
        payments = Payment.objects.filter(status__in=['Paid', 'Partial'], payment_date__gte=date_from, payment_date__lte=date_to).select_related('tenant', 'stall', 'collected_by')
        total_collected = payments.aggregate(total=Sum('amount_paid'))['total'] or 0
        template = 'reports/pdf/daily.html'
        context = {'payments': payments, 'date_from': date_from, 'date_to': date_to, 'total_collected': total_collected, 'report_type': 'Weekly Collection Report'}
    elif report_type == 'annual':
        payments = Payment.objects.filter(status__in=['Paid', 'Partial'], payment_date__year=year).select_related('tenant', 'stall', 'collected_by')
        total_collected = payments.aggregate(total=Sum('amount_paid'))['total'] or 0
        template = 'reports/pdf/daily.html'
        context = {'payments': payments, 'date_from': f'{year}-01-01', 'date_to': f'{year}-12-31', 'total_collected': total_collected, 'report_type': f'Annual Collection Report {year}'}
    elif report_type == 'section':
        from core.models import MarketSection
        sections = MarketSection.objects.annotate(stall_count=__import__('django.db.models', fromlist=['Count']).Count('stalls'))
        section_data = []
        for section in sections:
            section_stalls = Stall.objects.filter(section=section)
            occupied = section_stalls.filter(status='Occupied').count()
            vacant = section_stalls.filter(status='Vacant').count()
            billings = Billing.objects.filter(stall__section=section, billing_month=month, billing_year=year)
            total_billed = billings.aggregate(total=Sum('total_due'))['total'] or 0
            total_collected = billings.aggregate(total=Sum('amount_paid'))['total'] or 0
            total_balance = billings.aggregate(total=Sum('balance'))['total'] or 0
            section_data.append({'section': section, 'total_stalls': section_stalls.count(), 'occupied': occupied, 'vacant': vacant, 'total_billed': total_billed, 'total_collected': total_collected, 'total_balance': total_balance})
        template = 'reports/pdf/section.html'
        context = {'section_data': section_data, 'month': month, 'year': year, 'month_name': __import__('calendar').month_name[month], 'report_type': 'Collection by Section'}
    elif report_type == 'collector':
        from core.models import UserProfile
        collectors = UserProfile.objects.filter(role__in=['collector', 'cashier', 'admin']).select_related('user')
        collector_data = []
        for profile in collectors:
            payments = Payment.objects.filter(collected_by=profile.user, status__in=['Paid', 'Partial'], payment_date__gte=date_from, payment_date__lte=date_to)
            total = payments.aggregate(total=Sum('amount_paid'))['total'] or 0
            count = payments.count()
            collector_data.append({'collector': profile.user, 'total': total, 'count': count})
        grand_total = sum(d['total'] for d in collector_data)
        template = 'reports/pdf/collector.html'
        context = {'collector_data': collector_data, 'date_from': date_from, 'date_to': date_to, 'grand_total': grand_total, 'report_type': 'Collection by Collector'}
    elif report_type in ['paid', 'paid_tenants']:
        billings = Billing.objects.filter(billing_month=month, billing_year=year, status='Paid').select_related('tenant', 'stall')
        total_billed = billings.aggregate(total=Sum('total_due'))['total'] or 0
        total_collected = billings.aggregate(total=Sum('amount_paid'))['total'] or 0
        template = 'reports/pdf/monthly.html'
        context = {'billings': billings, 'month': month, 'year': year, 'month_name': __import__('calendar').month_name[month], 'total_billed': total_billed, 'total_collected': total_collected, 'report_type': 'Paid Tenants Report'}
    elif report_type in ['unpaid', 'unpaid_tenants']:
        billings = Billing.objects.filter(billing_month=month, billing_year=year).filter(Q(status='Unpaid') | Q(status='Partial') | Q(status='Overdue')).select_related('tenant', 'stall')
        total_billed = billings.aggregate(total=Sum('total_due'))['total'] or 0
        total_collected = billings.aggregate(total=Sum('amount_paid'))['total'] or 0
        template = 'reports/pdf/monthly.html'
        context = {'billings': billings, 'month': month, 'year': year, 'month_name': __import__('calendar').month_name[month], 'total_billed': total_billed, 'total_collected': total_collected, 'report_type': 'Unpaid Tenants Report'}
    elif report_type == 'overdue':
        billings = Billing.objects.filter(status='Overdue').select_related('tenant', 'stall')
        total_billed = billings.aggregate(total=Sum('total_due'))['total'] or 0
        total_collected = billings.aggregate(total=Sum('amount_paid'))['total'] or 0
        template = 'reports/pdf/monthly.html'
        context = {'billings': billings, 'month': month, 'year': year, 'month_name': __import__('calendar').month_name[month], 'total_billed': total_billed, 'total_collected': total_collected, 'report_type': 'Overdue Accounts Report'}
    elif report_type in ['occupancy', 'vacant']:
        qs = Stall.objects.select_related('section', 'stall_type').all().order_by('section__name', 'stall_number')
        if report_type == 'vacant':
            qs = qs.filter(status='Vacant')
        template = 'reports/pdf/occupancy.html'
        context = {'stalls': qs, 'total': qs.count(), 'report_type': 'Vacant Stall Report' if report_type == 'vacant' else 'Stall Occupancy Report'}
    elif report_type == 'ledger':
        tenant_id = request.GET.get('tenant')
        entries = TenantLedger.objects.select_related('tenant', 'billing', 'payment').order_by('transaction_date')
        if tenant_id:
            entries = entries.filter(tenant_id=tenant_id)
        if date_from:
            entries = entries.filter(transaction_date__gte=date_from)
        if date_to:
            entries = entries.filter(transaction_date__lte=date_to)
        total_debit = entries.aggregate(total=Sum('debit'))['total'] or 0
        total_credit = entries.aggregate(total=Sum('credit'))['total'] or 0
        template = 'reports/pdf/ledger.html'
        context = {'entries': entries, 'date_from': date_from, 'date_to': date_to, 'total_debit': total_debit, 'total_credit': total_credit, 'report_type': 'Tenant Ledger Report'}
    else:
        billings = Billing.objects.filter(billing_month=month, billing_year=year).select_related('tenant', 'stall')
        total_billed = billings.aggregate(total=Sum('total_due'))['total'] or 0
        total_collected = billings.aggregate(total=Sum('amount_paid'))['total'] or 0
        template = 'reports/pdf/monthly.html'
        context = {'billings': billings, 'month': month, 'year': year, 'month_name': __import__('calendar').month_name[month], 'total_billed': total_billed, 'total_collected': total_collected, 'report_type': 'Collection Report'}

    html_string = render(request, template, context).content
    # weasyprint expects str, decode if bytes
    if isinstance(html_string, bytes):
        html_string = html_string.decode('utf-8')
    pdf_file = weasyprint.HTML(string=html_string).write_pdf()
    response = HttpResponse(pdf_file, content_type='application/pdf')
    response['Content-Disposition'] = f'attachment; filename="{report_type}_report_{date.today()}.pdf"'
    return response


@staff_required
def report_export_excel(request):
    report_type = request.GET.get('report', 'monthly')
    month, year = get_month_year(request)
    date_from, date_to = get_date_range(request)

    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        messages.error(request, 'openpyxl is not installed. Excel export is unavailable.')
        return render(request, 'reports/export_excel.html', {'error': True})

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f'{report_type.title()} Report'

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='115E59', end_color='115E59', fill_type='solid')

    def style_header(headers):
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')

    if report_type == 'monthly':
        billings = Billing.objects.filter(billing_month=month, billing_year=year).select_related('tenant', 'stall', 'stall__section')
        headers = ['Tenant ID', 'Tenant Name', 'Stall', 'Section', 'Rental', 'Penalty', 'Total Due', 'Amount Paid', 'Balance', 'Status']
        style_header(headers)
        for row, billing in enumerate(billings, 2):
            ws.cell(row=row, column=1, value=billing.tenant.tenant_id)
            ws.cell(row=row, column=2, value=billing.tenant.full_name)
            ws.cell(row=row, column=3, value=billing.stall.stall_number)
            ws.cell(row=row, column=4, value=billing.stall.section.name)
            ws.cell(row=row, column=5, value=float(billing.rental_amount))
            ws.cell(row=row, column=6, value=float(billing.penalty_amount))
            ws.cell(row=row, column=7, value=float(billing.total_due))
            ws.cell(row=row, column=8, value=float(billing.amount_paid))
            ws.cell(row=row, column=9, value=float(billing.balance))
            ws.cell(row=row, column=10, value=billing.get_status_display())
    elif report_type in ['daily', 'weekly']:
        payments = Payment.objects.filter(status__in=['Paid', 'Partial'], payment_date__gte=date_from, payment_date__lte=date_to).select_related('tenant', 'stall', 'collected_by')
        headers = ['Receipt#', 'Date', 'Tenant', 'Stall', 'Amount Paid', 'Method', 'Collected By', 'Status']
        style_header(headers)
        for row, payment in enumerate(payments, 2):
            ws.cell(row=row, column=1, value=payment.receipt_number)
            ws.cell(row=row, column=2, value=payment.payment_date.isoformat())
            ws.cell(row=row, column=3, value=payment.tenant.full_name)
            ws.cell(row=row, column=4, value=payment.stall.stall_number)
            ws.cell(row=row, column=5, value=float(payment.amount_paid))
            ws.cell(row=row, column=6, value=payment.get_payment_method_display())
            ws.cell(row=row, column=7, value=payment.collected_by.get_full_name() if payment.collected_by else '')
            ws.cell(row=row, column=8, value=payment.get_status_display())
    elif report_type == 'annual':
        headers = ['Month', 'Transactions', 'Total Collected', '% of Grand']
        style_header(headers)
        monthly_data = []
        for m in range(1, 13):
            payments = Payment.objects.filter(status__in=['Paid', 'Partial'], payment_date__year=year, payment_date__month=m)
            total = payments.aggregate(total=Sum('amount_paid'))['total'] or 0
            count = payments.count()
            monthly_data.append((m, count, total))
        grand = sum(t for _, _, t in monthly_data)
        for idx, (m, count, total) in enumerate(monthly_data, 2):
            ws.cell(row=idx, column=1, value=__import__('calendar').month_name[m])
            ws.cell(row=idx, column=2, value=count)
            ws.cell(row=idx, column=3, value=float(total))
            ws.cell(row=idx, column=4, value=round(float(total)/float(grand)*100,2) if grand else 0)
    elif report_type == 'section':
        from core.models import MarketSection
        sections = MarketSection.objects.all()
        headers = ['Section', 'Stalls', 'Occupied', 'Vacant', 'Billed', 'Collected', 'Balance']
        style_header(headers)
        for row, section in enumerate(sections, 2):
            section_stalls = Stall.objects.filter(section=section)
            occupied = section_stalls.filter(status='Occupied').count()
            vacant = section_stalls.filter(status='Vacant').count()
            billings = Billing.objects.filter(stall__section=section, billing_month=month, billing_year=year)
            total_billed = billings.aggregate(total=Sum('total_due'))['total'] or 0
            total_collected = billings.aggregate(total=Sum('amount_paid'))['total'] or 0
            total_balance = billings.aggregate(total=Sum('balance'))['total'] or 0
            ws.cell(row=row, column=1, value=section.name)
            ws.cell(row=row, column=2, value=section_stalls.count())
            ws.cell(row=row, column=3, value=occupied)
            ws.cell(row=row, column=4, value=vacant)
            ws.cell(row=row, column=5, value=float(total_billed))
            ws.cell(row=row, column=6, value=float(total_collected))
            ws.cell(row=row, column=7, value=float(total_balance))
    elif report_type == 'collector':
        from core.models import UserProfile
        collectors = UserProfile.objects.filter(role__in=['collector', 'cashier', 'admin']).select_related('user')
        headers = ['Collector', 'Transactions', 'Total Collected']
        style_header(headers)
        for row, profile in enumerate(collectors, 2):
            payments = Payment.objects.filter(collected_by=profile.user, status__in=['Paid', 'Partial'], payment_date__gte=date_from, payment_date__lte=date_to)
            total = payments.aggregate(total=Sum('amount_paid'))['total'] or 0
            ws.cell(row=row, column=1, value=profile.user.get_full_name() or profile.user.username)
            ws.cell(row=row, column=2, value=payments.count())
            ws.cell(row=row, column=3, value=float(total))
    elif report_type in ['paid', 'paid_tenants']:
        billings = Billing.objects.filter(billing_month=month, billing_year=year, status='Paid').select_related('tenant', 'stall')
        headers = ['Tenant ID', 'Tenant Name', 'Stall', 'Section', 'Rental', 'Penalty', 'Total Due', 'Amount Paid', 'Balance', 'Status']
        style_header(headers)
        for row, billing in enumerate(billings, 2):
            ws.cell(row=row, column=1, value=billing.tenant.tenant_id)
            ws.cell(row=row, column=2, value=billing.tenant.full_name)
            ws.cell(row=row, column=3, value=billing.stall.stall_number)
            ws.cell(row=row, column=4, value=billing.stall.section.name)
            ws.cell(row=row, column=5, value=float(billing.rental_amount))
            ws.cell(row=row, column=6, value=float(billing.penalty_amount))
            ws.cell(row=row, column=7, value=float(billing.total_due))
            ws.cell(row=row, column=8, value=float(billing.amount_paid))
            ws.cell(row=row, column=9, value=float(billing.balance))
            ws.cell(row=row, column=10, value=billing.get_status_display())
    elif report_type in ['unpaid', 'unpaid_tenants']:
        billings = Billing.objects.filter(billing_month=month, billing_year=year).filter(Q(status='Unpaid') | Q(status='Partial') | Q(status='Overdue')).select_related('tenant', 'stall')
        headers = ['Tenant ID', 'Tenant Name', 'Stall', 'Section', 'Rental', 'Penalty', 'Total Due', 'Amount Paid', 'Balance', 'Status']
        style_header(headers)
        for row, billing in enumerate(billings, 2):
            ws.cell(row=row, column=1, value=billing.tenant.tenant_id)
            ws.cell(row=row, column=2, value=billing.tenant.full_name)
            ws.cell(row=row, column=3, value=billing.stall.stall_number)
            ws.cell(row=row, column=4, value=billing.stall.section.name)
            ws.cell(row=row, column=5, value=float(billing.rental_amount))
            ws.cell(row=row, column=6, value=float(billing.penalty_amount))
            ws.cell(row=row, column=7, value=float(billing.total_due))
            ws.cell(row=row, column=8, value=float(billing.amount_paid))
            ws.cell(row=row, column=9, value=float(billing.balance))
            ws.cell(row=row, column=10, value=billing.get_status_display())
    elif report_type == 'overdue':
        billings = Billing.objects.filter(status='Overdue').select_related('tenant', 'stall')
        headers = ['Tenant ID', 'Tenant Name', 'Stall', 'Section', 'Rental', 'Penalty', 'Total Due', 'Amount Paid', 'Balance', 'Due Date', 'Status']
        style_header(headers)
        for row, billing in enumerate(billings, 2):
            ws.cell(row=row, column=1, value=billing.tenant.tenant_id)
            ws.cell(row=row, column=2, value=billing.tenant.full_name)
            ws.cell(row=row, column=3, value=billing.stall.stall_number)
            ws.cell(row=row, column=4, value=billing.stall.section.name)
            ws.cell(row=row, column=5, value=float(billing.rental_amount))
            ws.cell(row=row, column=6, value=float(billing.penalty_amount))
            ws.cell(row=row, column=7, value=float(billing.total_due))
            ws.cell(row=row, column=8, value=float(billing.amount_paid))
            ws.cell(row=row, column=9, value=float(billing.balance))
            ws.cell(row=row, column=10, value=billing.due_date.isoformat())
            ws.cell(row=row, column=11, value=billing.get_status_display())
    elif report_type in ['occupancy', 'vacant']:
        qs = Stall.objects.select_related('section', 'stall_type').all().order_by('section__name', 'stall_number')
        if report_type == 'vacant':
            qs = qs.filter(status='Vacant')
        headers = ['Stall #', 'Section', 'Type', 'Size', 'Monthly Rate', 'Status', 'Tenant']
        style_header(headers)
        for row, stall in enumerate(qs, 2):
            tenant_name = ''
            active_contract = stall.contracts.filter(status='Active').first()
            if active_contract:
                tenant_name = active_contract.tenant.full_name
            ws.cell(row=row, column=1, value=stall.stall_number)
            ws.cell(row=row, column=2, value=stall.section.name)
            ws.cell(row=row, column=3, value=stall.stall_type.name if stall.stall_type else '--')
            ws.cell(row=row, column=4, value=float(stall.size) if stall.size else 0)
            ws.cell(row=row, column=5, value=float(stall.monthly_rate))
            ws.cell(row=row, column=6, value=stall.status)
            ws.cell(row=row, column=7, value=tenant_name)
    elif report_type == 'ledger':
        tenant_id = request.GET.get('tenant')
        entries = TenantLedger.objects.select_related('tenant', 'billing', 'payment').order_by('transaction_date')
        if tenant_id:
            entries = entries.filter(tenant_id=tenant_id)
        if date_from:
            entries = entries.filter(transaction_date__gte=date_from)
        if date_to:
            entries = entries.filter(transaction_date__lte=date_to)
        headers = ['Date', 'Tenant', 'Description', 'Debit', 'Credit', 'Balance']
        style_header(headers)
        for row, entry in enumerate(entries, 2):
            ws.cell(row=row, column=1, value=entry.transaction_date.isoformat())
            ws.cell(row=row, column=2, value=entry.tenant.full_name)
            ws.cell(row=row, column=3, value=entry.description)
            ws.cell(row=row, column=4, value=float(entry.debit))
            ws.cell(row=row, column=5, value=float(entry.credit))
            ws.cell(row=row, column=6, value=float(entry.balance))
    else:
        payments = Payment.objects.filter(status__in=['Paid', 'Partial'], payment_date__month=month, payment_date__year=year).select_related('tenant', 'stall', 'collected_by')
        headers = ['Receipt#', 'Date', 'Tenant', 'Stall', 'Amount Paid', 'Method', 'Collected By', 'Status']
        style_header(headers)
        for row, payment in enumerate(payments, 2):
            ws.cell(row=row, column=1, value=payment.receipt_number)
            ws.cell(row=row, column=2, value=payment.payment_date.isoformat())
            ws.cell(row=row, column=3, value=payment.tenant.full_name)
            ws.cell(row=row, column=4, value=payment.stall.stall_number)
            ws.cell(row=row, column=5, value=float(payment.amount_paid))
            ws.cell(row=row, column=6, value=payment.get_payment_method_display())
            ws.cell(row=row, column=7, value=payment.collected_by.get_full_name() if payment.collected_by else '')
            ws.cell(row=row, column=8, value=payment.get_status_display())

    for column in ws.columns:
        max_length = 0
        col_letter = column[0].column_letter
        for cell in column:
            if cell.value:
                max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max_length + 3

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{report_type}_report_{date.today()}.xlsx"'
    wb.save(response)
    return response
